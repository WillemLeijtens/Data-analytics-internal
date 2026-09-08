"""Douglas — iCube "Advanced Sell Out NET": twee maandrapporten.

Douglas levert per maand twee exports die dezelfde netto-omzet op twee
manieren verdelen:

  * SKU-rapport    Country x Channel x Brand x EAN x Month
  * Winkelrapport  Country x Channel x Brand x Store x Month

Beide hebben vier bedragkolommen — deze maand, dezelfde maand vorig jaar,
YTD-cumulatief en YTD-cumulatief vorig jaar — en géén volume. Beide tellen
op tot precies hetzelfde totaal, dus naast elkaar in de feitentabel zouden
ze de omzet verdubbelen. Daarom krijgt elke regel een `niveau` ('artikel' of
'winkel', zie migratie 023): de analyses lezen er één.

Wat de kolommen verder opleveren:

  * De LY-kolom is dezelfde maand vorig jaar. Die gaat als eigen regel mee
    (periode een jaar terug), zodat de jaarvergelijking uit één bestand
    werkt. Een later geladen écht bestand van vorig jaar heeft dezelfde
    sleutel en vervangt die regels gewoon.
  * De Cum-kolommen gaan niet als feit mee (er is geen maandverdeling van
    te maken), maar wel als totaal terug naar de importer: staat er meer
    cumulatief dan deze maand terwijl er nog geen eerdere maanden geladen
    zijn, dan ontbreken er bestanden en hoort dat gemeld te worden.

Een regel zonder omzet deze maand maar mét cumulatief (het artikel of de
winkel deed dit jaar wél mee) wordt een 0-regel: gemeten, niets verkocht —
dezelfde keuze als bij Etos, en precies het signaal dat de stille-winkels-
en delist-logica nodig heeft. Een regel met alleen vorig jaar (geen omzet,
geen cumulatief) levert alleen de LY-regel op.

Kanaal wordt de formule: Brick & Mortar -> FYSIEK, E-Commerce -> ONLINE. De
webshop staat in het winkelrapport als pseudo-winkel (105 NL, 598 BE); het
dashboard houdt online-formules buiten "per winkel" via het profielveld
`online_banners`.
"""

from __future__ import annotations

import io
import re

from .cellen import als_identifier
from .periods import parse_period, period_number, period_year

LANDEN = {"netherlands": "NL", "belgium": "BE"}
KANALEN = {"brick & mortar": "FYSIEK", "e-commerce": "ONLINE"}

# De koplabels zoals iCube ze schrijft. Samengevoegde blokken (International
# Article, Store II) hebben één label in de eerste cel en lege cellen erna.
KOP_LAND = "country ii"
KOP_KANAAL = "channel ii"
KOP_MERK = "international brand"
KOP_ARTIKEL = "international article"
KOP_EAN = "ean"
KOP_WINKEL = "store ii"
KOP_MAAND = "month"
KOP_OMZET = "sell-out net eur"
KOP_OMZET_LY = "sell-out net ly eur"
KOP_CUM = "sell-out net cy cum eur"
KOP_CUM_LY = "sell-out net cy cum ly eur"

_WIT = re.compile(r"\s+")


def _norm(v) -> str:
    return _WIT.sub(" ", str(v)).strip() if v is not None else ""


def _num(v) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        raise ValueError(f"bedrag {v!r} is geen getal")


def _kop(rij) -> dict[str, int]:
    """{label (lowercase): kolomindex} van de koprij; lege cellen slaan we over."""
    return {_norm(c).lower(): i for i, c in enumerate(rij) if _norm(c)}


def _kolomblok(kop: dict[str, int], start_label: str, volgorde: list[str]) -> range | None:
    """De kolommen van een samengevoegd blok: vanaf zijn label tot het
    eerstvolgende gelabelde kolom."""
    if start_label not in kop:
        return None
    start = kop[start_label]
    volgende = min((i for i in volgorde if i > start), default=start + 1)
    return range(start, volgende)


def _rapporttype(kop: dict[str, int]) -> str | None:
    basis = {KOP_LAND, KOP_KANAAL, KOP_MERK, KOP_MAAND, KOP_OMZET}
    if not basis <= set(kop):
        return None
    if KOP_EAN in kop and KOP_ARTIKEL in kop:
        return "artikel"
    if KOP_WINKEL in kop:
        return "winkel"
    return None


def content_matches(content: bytes) -> bool:
    """Herkenning op structuur, voor een hernoemd bestand. Nooit een exception."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        for rij in ws.iter_rows(min_row=1, max_row=1, values_only=True):
            return _rapporttype(_kop(rij)) is not None
        return False
    except Exception:  # noqa: BLE001 - onleesbaar bestand matcht gewoon niet
        return False


def _vorig_jaar(periode: str) -> str:
    return f"{period_year(periode) - 1}-{period_number(periode):02d}"


def parse_workbook(content: bytes) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rijen = ws.iter_rows(min_row=1, values_only=True)
    try:
        koprij = next(rijen)
    except StopIteration:
        raise ValueError("leeg werkblad")
    kop = _kop(koprij)
    soort = _rapporttype(kop)
    if soort is None:
        raise ValueError(
            "koprij niet herkend als Douglas iCube-export (verwacht Country II, "
            "Channel II, International Brand, Month, Sell-Out NET EUR en EAN of Store II)")
    gelabeld = sorted(kop.values())
    if soort == "artikel":
        naamblok = _kolomblok(kop, KOP_ARTIKEL, gelabeld)
        col_id = kop[KOP_EAN]
    else:
        naamblok = _kolomblok(kop, KOP_WINKEL, gelabeld)
        col_id = naamblok.start
    for vereist in (KOP_OMZET_LY, KOP_CUM, KOP_CUM_LY):
        if vereist not in kop:
            raise ValueError(f"kolom '{vereist}' ontbreekt in de koprij")

    facts: list[dict] = []
    ly_regels = 0
    maand_totaal = cum_totaal = 0.0
    periodes: set[str] = set()
    for nr, r in enumerate(rijen, 2):
        if not any(_norm(c) for c in r):
            continue
        cel = lambda i: r[i] if i < len(r) else None  # noqa: E731

        land_raw = _norm(cel(kop[KOP_LAND]))
        kanaal_raw = _norm(cel(kop[KOP_KANAAL]))
        land = LANDEN.get(land_raw.lower())
        banner = KANALEN.get(kanaal_raw.lower())
        # Fail closed: een onbekend land of kanaal zou als eigen scope in
        # instellingen, dekking en filters belanden. Beter één duidelijke
        # melding dan een stille nieuwe categorie.
        if not land:
            raise ValueError(f"rij {nr}: onbekend land {land_raw!r} (verwacht Netherlands/Belgium)")
        if not banner:
            raise ValueError(f"rij {nr}: onbekend kanaal {kanaal_raw!r} "
                             "(verwacht Brick & Mortar / E-Commerce)")
        merk = _norm(cel(kop[KOP_MERK])).upper()
        if not merk:
            raise ValueError(f"rij {nr}: merk ontbreekt")
        periode = parse_period(cel(kop[KOP_MAAND]), "mmm-yy")

        omzet = _num(cel(kop[KOP_OMZET]))
        omzet_ly = _num(cel(kop[KOP_OMZET_LY]))
        cum = _num(cel(kop[KOP_CUM]))

        if soort == "artikel":
            ean = als_identifier(cel(col_id))
            if not ean:
                # Zonder EAN valt de regel buiten elke artikelanalyse én zou
                # de korrelwissel-regel van de importer hem als
                # "artikelloos" kunnen wegvegen. Niet stil overslaan.
                raise ValueError(f"rij {nr}: EAN ontbreekt")
            delen = [_norm(cel(i)) for i in naamblok]
            delen = [d for d in delen if d]
            naam = " ".join(delen) or None
            categorie = delen[0] if delen else None
            basis = {"merk": merk, "land": land, "banner": banner,
                     "winkel_id": None, "winkel_naam": None,
                     "artikel_ean": ean, "artikel_naam": naam, "categorie": categorie}
        else:
            winkel_id = als_identifier(cel(col_id))
            if not winkel_id:
                raise ValueError(f"rij {nr}: winkelnummer ontbreekt")
            rest = [_norm(cel(i)) for i in naamblok][1:]
            stad = rest[0] if rest else ""
            adres = " ".join(d for d in rest[1:] if d)
            naam = ", ".join(d for d in (stad, adres) if d) or None
            basis = {"merk": merk, "land": land, "banner": banner,
                     "winkel_id": winkel_id, "winkel_naam": naam,
                     "artikel_ean": None, "artikel_naam": None, "categorie": None}

        # Deze maand: een bedrag, of een 0-regel als er dit jaar wél iets
        # cumulatief staat (gemeten, niets verkocht). Alleen-vorig-jaar
        # levert geen regel voor dit jaar op.
        if omzet is not None or cum is not None:
            facts.append({**basis, "periode": periode, "volume": 0,
                          "omzet": float(omzet or 0.0)})
            periodes.add(periode)
            maand_totaal += omzet or 0.0
            cum_totaal += cum or 0.0
        if omzet_ly is not None:
            facts.append({**basis, "periode": _vorig_jaar(periode), "volume": 0,
                          "omzet": float(omzet_ly)})
            ly_regels += 1

    if not facts:
        raise ValueError("geen datarijen met cijfers gevonden")

    # Fail-closed dubbelcheck binnen het bestand, zoals de andere parsers.
    keys = [(f["merk"], f["land"], f["banner"], f["artikel_ean"] or f["winkel_id"], f["periode"])
            for f in facts]
    if len(keys) != len(set(keys)):
        dubbel = sorted({k for k in keys if keys.count(k) > 1})[:5]
        raise ValueError(
            f"{len(keys) - len(set(keys))} dubbele combinatie(s) in het bestand "
            f"(o.a. {dubbel}); import afgebroken om dubbeltelling te voorkomen")

    warnings: list[str] = []
    if ly_regels:
        vorig = sorted({_vorig_jaar(p) for p in periodes})
        warnings.append(
            f"inclusief {ly_regels} regel(s) voor {', '.join(vorig)} uit de kolom "
            "'Sell-Out NET LY EUR' (dezelfde maand vorig jaar)")

    return {"facts": facts, "periode_type": "maand",
            # Alleen de maand van het bestand zelf: de LY-regels zijn een
            # afgeleide, en "2025-08 t/m 2026-08 (2)" in de importstatus zou
            # doen alsof er twaalf maanden geleverd zijn.
            "periodes": sorted(periodes),
            "warnings": warnings, "niveau": soort,
            # Voor de importer: is er meer cumulatief dan deze maand, dan
            # bestaan er eerdere maandbestanden die nog niet geladen zijn.
            "cumulatief": {"maand": round(maand_totaal, 2), "totaal": round(cum_totaal, 2)}}
